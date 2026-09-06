"""Structured office operations used by both Python and desktop agents."""

from __future__ import annotations

import hashlib
import json
import os
import tempfile
from collections.abc import Generator, Iterator
from io import BytesIO
from pathlib import Path
from typing import Any, Literal
from zipfile import ZipFile

from pydantic import BaseModel, ConfigDict, Field, StrictBool, StrictFloat, StrictInt, StrictStr

from sztu_code.core.documents import MAX_FILE_BYTES, DocumentError
from sztu_code.core.tools.workspace import resolve_workspace_path

CellValue = StrictStr | StrictInt | StrictFloat | StrictBool | None


class Params(BaseModel):
    model_config = ConfigDict(extra="forbid")
    path: str = Field(min_length=1, description="Output or input path relative to the workspace")


class ReadDocumentParams(Params):
    offset: int = Field(default=0, ge=0, le=1000000, description="Block offset; use next_offset")
    limit: int = Field(default=40, ge=1, le=100)
    sheet: str | None = Field(default=None, description="XLSX: exact worksheet name")


class WordBlock(BaseModel):
    model_config = ConfigDict(extra="forbid")
    kind: Literal["heading", "paragraph", "bullet", "table"]
    text: str = Field(default="", max_length=50000)
    level: int = Field(default=1, ge=1, le=6)
    rows: list[list[str]] = Field(default_factory=list, max_length=1000)


class SheetSpec(BaseModel):
    model_config = ConfigDict(extra="forbid")
    name: str = Field(min_length=1, max_length=31)
    rows: list[list[CellValue]] = Field(min_length=1, max_length=10000)


class SlideSpec(BaseModel):
    model_config = ConfigDict(extra="forbid")
    title: str = Field(max_length=120)
    bullets: list[str] = Field(default_factory=list, max_length=8)
    notes: str = Field(default="", max_length=20000)


class CreateDocumentParams(Params):
    title: str = Field(default="", max_length=200)
    blocks: list[WordBlock] = Field(
        default_factory=list, max_length=1000, description="DOCX body in document order"
    )
    sheets: list[SheetSpec] = Field(
        default_factory=list,
        max_length=100,
        description="XLSX sheets; strings starting with = are formulas",
    )
    slides: list[SlideSpec] = Field(
        default_factory=list,
        max_length=100,
        description="PPTX title and bullet slides, with speaker notes",
    )
    overwrite: bool = False


class EditOperation(BaseModel):
    model_config = ConfigDict(extra="forbid")
    kind: Literal["replace_text", "set_cell"]
    find: str = Field(default="", max_length=10000)
    replace: str = Field(default="", max_length=50000)
    expected_matches: int = Field(
        default=1, ge=1, le=10000, description="replace_text: exact expected count across source"
    )
    sheet: str = ""
    cell: str = Field(default="", description="set_cell: Excel coordinate, e.g. B3")
    value: CellValue = None


class EditDocumentParams(Params):
    source: str = Field(
        min_length=1, description="Existing source document; saved to path as a copy"
    )
    source_sha256: str | None = Field(default=None, pattern="^[a-f0-9]{64}$")
    operations: list[EditOperation] = Field(min_length=1, max_length=500)
    overwrite: bool = False


MODELS: dict[
    str, type[ReadDocumentParams] | type[CreateDocumentParams] | type[EditDocumentParams]
] = {
    "read_document": ReadDocumentParams,
    "create_document": CreateDocumentParams,
    "edit_document": EditDocumentParams,
}
DESCRIPTIONS = {
    "read_document": (
        "读取 PDF、Word(DOCX)、Excel(XLSX)、PPT(PPTX) 的结构化内容，"
        "返回可引用的页码、段落、工作表/单元格、幻灯片和备注位置。"
        "长资料用 next_offset 翻页；不要把第一页当成全文。"
    ),
    "create_document": (
        "生成真正可编辑的 DOCX、XLSX 或 PPTX 文件。分别提供 blocks、sheets 或 slides；"
        "默认不覆盖文件。XLSX 保留公式表达式但不计算。"
        "保存后重开校验结构，未进行视觉排版校验。"
    ),
    "edit_document": (
        "修改办公文件并另存副本：DOCX/PPTX 支持跨文本运行的精确替换，"
        "XLSX 支持按工作表和单元格写值/公式。先 read_document，"
        "再使用 source_sha256 防止修改过期内容；替换次数须匹配 expected_matches。"
    ),
}


def office_contract() -> list[dict[str, Any]]:
    return [
        {
            "name": name,
            "description": DESCRIPTIONS[name],
            "permission": "read_only" if name == "read_document" else "workspace_write",
            "schema": model.model_json_schema(),
        }
        for name, model in MODELS.items()
    ]


def _load(path: Path) -> tuple[bytes, str]:
    if path.stat().st_size > MAX_FILE_BYTES:
        raise DocumentError("资料文件超过 20MB 限制")
    with path.open("rb") as stream:
        data = stream.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError("资料文件超过 20MB 限制")
    format = "pdf" if data.startswith(b"%PDF-") else path.suffix.lower()[1:]
    if format not in {"pdf", "docx", "xlsx", "pptx"}:
        raise DocumentError("支持 PDF、DOCX、XLSX、PPTX；旧版 DOC/XLS/PPT 请先转换格式")
    if format != "pdf":
        with ZipFile(BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise DocumentError("文档解压后超过 100MB 限制")
    return data, format


def _text_blocks(location: str, text: str, **metadata: Any) -> Iterator[dict[str, Any]]:
    # Split large paragraphs/cells instead of permanently discarding their tail.
    for start in range(0, len(text), 2000):
        yield {
            "location": location,
            "text_offset": start,
            "text": text[start : start + 2000],
            **metadata,
        }


def _word_paragraphs(container: Any, prefix: str = "body") -> Iterator[tuple[str, Any]]:
    from docx.table import Table

    for index, block in enumerate(container.iter_inner_content(), 1):
        location = f"{prefix}/block:{index}"
        if isinstance(block, Table):
            seen = set()
            for row_index, row in enumerate(block.rows, 1):
                for col_index, cell in enumerate(row.cells, 1):
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    yield from _word_paragraphs(cell, f"{location}/row:{row_index}/col:{col_index}")
        else:
            yield location, block


def _word_all_paragraphs(doc: Any) -> Iterator[tuple[str, Any]]:
    yield from _word_paragraphs(doc)
    seen = set()
    for index, section in enumerate(doc.sections, 1):
        for name in (
            "header",
            "footer",
            "first_page_header",
            "first_page_footer",
            "even_page_header",
            "even_page_footer",
        ):
            part = getattr(section, name)
            # Reading an absent header must not create one.
            if part.is_linked_to_previous or part.part.partname in seen:
                continue
            seen.add(part.part.partname)
            yield from _word_paragraphs(part, f"section:{index}/{name}")


def _slide_paragraphs(shapes: Any, prefix: str) -> Iterator[tuple[str, Any]]:
    from pptx.shapes.group import GroupShape

    for shape in shapes:
        location = f"{prefix}/shape:{shape.shape_id}"
        if shape.has_text_frame:
            for index, paragraph in enumerate(shape.text_frame.paragraphs, 1):
                yield f"{location}/paragraph:{index}", paragraph
        if shape.has_table:
            for r, row in enumerate(shape.table.rows, 1):
                for c, cell in enumerate(row.cells, 1):
                    for index, paragraph in enumerate(cell.text_frame.paragraphs, 1):
                        yield f"{location}/row:{r}/col:{c}/paragraph:{index}", paragraph
        if isinstance(shape, GroupShape):
            yield from _slide_paragraphs(shape.shapes, location)


def _ppt_all_paragraphs(doc: Any) -> Iterator[tuple[str, Any]]:
    for number, slide in enumerate(doc.slides, 1):
        yield from _slide_paragraphs(slide.shapes, f"slide:{number}")
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                for index, paragraph in enumerate(frame.paragraphs, 1):
                    yield f"slide:{number}/notes:{index}", paragraph


def _read_blocks(
    data: bytes, format: str, sheet: str | None, metadata: dict[str, Any]
) -> Generator[dict[str, Any]]:
    doc: Any
    if sheet is not None and format != "xlsx":
        raise DocumentError("sheet 仅适用于 XLSX")
    if format == "docx":
        from docx import Document

        doc = Document(BytesIO(data))
        metadata["title"] = doc.core_properties.title
        for location, paragraph in _word_all_paragraphs(doc):
            yield from _text_blocks(location, paragraph.text, style=paragraph.style.name)
    elif format == "pptx":
        from pptx import Presentation

        doc = Presentation(BytesIO(data))
        metadata["slide_count"] = len(doc.slides)
        for location, paragraph in _ppt_all_paragraphs(doc):
            yield from _text_blocks(location, paragraph.text)
    elif format == "xlsx":
        from openpyxl import load_workbook

        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            metadata["sheets"] = book.sheetnames
            metadata["formula_policy"] = "返回公式原文，不计算公式，不把公式当作计算结果。"
            if sheet is not None and sheet not in book.sheetnames:
                raise DocumentError(f"工作表不存在：{sheet}")
            for worksheet in [book[sheet]] if sheet else book.worksheets:
                for row in worksheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            yield from _text_blocks(
                                f"{worksheet.title}!{cell.coordinate}",
                                str(cell.value),
                                data_type=cell.data_type,
                                number_format=cell.number_format,
                            )
        finally:
            book.close()
    else:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentError("PDF 已加密，请先解锁")
        metadata["page_count"] = len(reader.pages)
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text().strip()
            yield from _text_blocks(
                f"page:{number}", text or "[本页没有文本层，可能需要 OCR]", needs_ocr=not bool(text)
            )


def read_document(path: Path, params: ReadDocumentParams) -> dict[str, Any]:
    data, format = _load(path)
    metadata: dict[str, Any] = {"format": format, "sha256": hashlib.sha256(data).hexdigest()}
    blocks = _read_blocks(data, format, params.sheet, metadata)
    selected: list[dict[str, Any]] = []
    next_offset = None
    size = 0
    try:
        for index, block in enumerate(blocks):
            if index < params.offset:
                continue
            block_size = len(json.dumps(block, ensure_ascii=False))
            if len(selected) >= params.limit or (selected and size + block_size > 32000):
                next_offset = index
                break
            selected.append({"index": index, **block})
            size += block_size
    finally:
        blocks.close()
    return {
        **metadata,
        "path": params.path,
        "blocks": selected,
        "next_offset": next_offset,
        "complete": next_offset is None,
        "limitations": "仅提取文本结构；不识别图片、图表图像，不保证版式还原。",
    }


def _save(doc: Any, target: Path, overwrite: bool) -> dict[str, Any]:
    if target.exists() and not overwrite:
        raise DocumentError("目标文件已存在，请另选文件名或明确设置 overwrite=true")
    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, filename = tempfile.mkstemp(suffix=target.suffix, dir=target.parent)
    os.close(descriptor)
    temporary = Path(filename)
    try:
        doc.save(temporary)
        data, format = _load(temporary)
        # Reopen the saved artifact with its native library before publishing it.
        if format == "docx":
            from docx import Document

            Document(BytesIO(data))
        elif format == "pptx":
            from pptx import Presentation

            Presentation(BytesIO(data))
        else:
            from openpyxl import load_workbook

            check = load_workbook(BytesIO(data), read_only=True, data_only=False)
            check.close()
        if overwrite:
            os.replace(temporary, target)
        else:
            # Atomic no-clobber publication, including competing writers.
            os.link(temporary, target)
        return {
            "format": format,
            "size": len(data),
            "sha256": hashlib.sha256(data).hexdigest(),
            "verification": "saved_and_reopened",
            "visual_verified": False,
            "limitations": "已校验文件结构；未渲染检查版式。Excel 公式需在办公软件中重新计算。",
        }
    finally:
        temporary.unlink(missing_ok=True)


def create_document(target: Path, params: CreateDocumentParams) -> dict[str, Any]:
    doc: Any
    format = target.suffix.lower()[1:]
    if format == "docx" and params.blocks and not params.sheets and not params.slides:
        from docx import Document
        from docx.shared import Pt as WordPt

        doc = Document()
        doc.core_properties.title = params.title
        doc.styles["Normal"].font.size = WordPt(11)
        if params.title:
            doc.add_heading(params.title, 0)
        for block in params.blocks:
            if block.kind == "heading":
                doc.add_heading(block.text, block.level)
            elif block.kind == "table":
                if (
                    not block.rows
                    or not block.rows[0]
                    or any(len(r) != len(block.rows[0]) for r in block.rows)
                ):
                    raise DocumentError("Word 表格必须为非空矩形")
                table = doc.add_table(rows=0, cols=len(block.rows[0]))
                table.style = "Table Grid"
                for row in block.rows:
                    for cell, value in zip(table.add_row().cells, row, strict=True):
                        cell.text = value
            else:
                doc.add_paragraph(block.text, "List Bullet" if block.kind == "bullet" else None)
    elif format == "xlsx" and params.sheets and not params.blocks and not params.slides:
        import re

        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        doc = Workbook()
        doc.remove(doc.active)
        names: set[str] = set()
        for spec in params.sheets:
            if re.search(r"[\\/*?:\[\]]", spec.name) or spec.name.casefold() in names:
                raise DocumentError("工作表名非法或重复")
            names.add(spec.name.casefold())
            sheet = doc.create_sheet(spec.name)
            for sheet_row in spec.rows:
                if len(sheet_row) > 256:
                    raise DocumentError("表格最多支持 256 列")
                if any(isinstance(value, str) and len(value) > 32767 for value in sheet_row):
                    raise DocumentError("单元格文本超过 Excel 的 32767 字符限制")
                sheet.append(sheet_row)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="24476A")
    elif format == "pptx" and params.slides and not params.blocks and not params.sheets:
        from pptx import Presentation
        from pptx.util import Inches, Pt

        doc = Presentation()
        doc.slide_width, doc.slide_height = Inches(13.333), Inches(7.5)
        doc.core_properties.title = params.title
        for slide_spec in params.slides:
            if any(len(text) > 200 for text in slide_spec.bullets):
                raise DocumentError("每条幻灯片正文最多 200 字符，请拆分内容")
            slide = doc.slides.add_slide(doc.slide_layouts[1])
            slide.shapes.title.text = slide_spec.title
            frame = slide.placeholders[1].text_frame
            frame.clear()
            for index, text in enumerate(slide_spec.bullets):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = text
                paragraph.font.size = Pt(22)
            if slide_spec.notes:
                slide.notes_slide.notes_text_frame.text = slide_spec.notes
    else:
        raise DocumentError(
            "请为 DOCX 提供 blocks、XLSX 提供 sheets、PPTX 提供 slides，且只提供对应格式内容"
        )
    return _save(doc, target, params.overwrite)


def _replace_runs(paragraph: Any, find: str, replace: str) -> int:
    runs = list(paragraph.runs)
    text = "".join(run.text for run in runs)
    count = text.count(find)
    positions = []
    cursor = 0
    while (start := text.find(find, cursor)) >= 0:
        positions.append(start)
        cursor = start + len(find)
    # Work backwards so earlier match offsets remain valid; preserve other runs.
    for start in reversed(positions):
        end = start + len(find)
        cursor = 0
        for run in runs:
            length = len(run.text)
            left, right = max(start - cursor, 0), min(end - cursor, length)
            if left < right:
                insertion = replace if cursor <= start < cursor + length else ""
                run.text = run.text[:left] + insertion + run.text[right:]
            cursor += length
    return count


def edit_document(source: Path, target: Path, params: EditDocumentParams) -> dict[str, Any]:
    doc: Any
    if source == target:
        raise DocumentError("请将修改结果另存为新文件，path 不能与 source 相同")
    data, format = _load(source)
    if params.source_sha256 and hashlib.sha256(data).hexdigest() != params.source_sha256:
        raise DocumentError("源文件已变化，请重新 read_document 后再修改")
    if format == "pdf" or target.suffix.lower() != f".{format}":
        raise DocumentError("只支持同格式 DOCX/XLSX/PPTX 副本编辑")
    if format == "docx":
        from docx import Document

        doc = Document(BytesIO(data))
        paragraphs = list(_word_all_paragraphs(doc))
    elif format == "pptx":
        from pptx import Presentation

        doc = Presentation(BytesIO(data))
        paragraphs = list(_ppt_all_paragraphs(doc))
    else:
        from openpyxl import load_workbook

        doc = load_workbook(BytesIO(data), data_only=False, keep_links=True)
        paragraphs = []
    for operation in params.operations:
        if operation.kind == "set_cell" and format == "xlsx":
            import re

            from openpyxl.utils.cell import coordinate_to_tuple

            if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", operation.cell):
                raise DocumentError("无效单元格地址，请使用 A1、B3 等格式")
            row, col = coordinate_to_tuple(operation.cell)
            if row > 1048576 or col > 16384 or operation.sheet not in doc.sheetnames:
                raise DocumentError("工作表不存在或单元格地址超限")
            if isinstance(operation.value, str) and len(operation.value) > 32767:
                raise DocumentError("单元格文本超过 32767 字符限制")
            doc[operation.sheet][operation.cell] = operation.value
        elif operation.kind == "replace_text" and format in {"docx", "pptx"}:
            if not operation.find:
                raise DocumentError("find 不得为空")
            matches = sum(
                "".join(r.text for r in p.runs).count(operation.find) for _, p in paragraphs
            )
            if matches != operation.expected_matches:
                raise DocumentError(
                    f"替换匹配 {matches} 处，与 expected_matches={operation.expected_matches} 不符"
                )
            for _, paragraph in paragraphs:
                _replace_runs(paragraph, operation.find, operation.replace)
        else:
            raise DocumentError("DOCX/PPTX 使用 replace_text，XLSX 使用 set_cell")
    return _save(doc, target, params.overwrite)


def execute_office_request(request: dict[str, Any]) -> dict[str, Any]:
    name = request["tool"]
    if name not in MODELS:
        raise DocumentError("未知办公工具")
    values = dict(request["params"])
    values.pop("description", None)
    params = MODELS[name].model_validate(values)
    root = Path(request["workspace"])
    target = resolve_workspace_path(root, params.path)
    if isinstance(params, ReadDocumentParams):
        return read_document(target, params)
    if isinstance(params, CreateDocumentParams):
        result = create_document(target, params)
    else:
        result = edit_document(resolve_workspace_path(root, params.source), target, params)
    return {"path": params.path, **result}
