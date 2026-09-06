"""Local document extraction, shared by read_file and the desktop helper.

Run this file directly to emit one UTF-8 JSON result. Keeping the entry point
standalone lets the desktop bundle it without the daemon or any LLM SDKs.
"""

from __future__ import annotations

import json
import sys
from collections.abc import Generator, Iterator
from dataclasses import asdict, dataclass
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING
from zipfile import ZipFile

if TYPE_CHECKING:
    from pptx.shapes.shapetree import GroupShapes, SlideShapes

SUPPORTED_FORMATS = {"pdf", "docx", "xlsx", "pptx"}
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 32 * 1024
MAX_PAGES = 200
TRUNCATION_NOTE = "\n[资料内容已截断：达到解析上限，可拆分文件后继续上传。]"


class DocumentError(ValueError):
    """A document cannot provide usable text."""


@dataclass
class ParsedDocument:
    text: str
    format: str
    truncated: bool = False


def _pdf_parts(data: bytes) -> Generator[str]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise DocumentError("PDF 已加密，请先解锁文件后重试。")
    found_text = False
    empty_pages = 0
    for index, page in enumerate(reader.pages):
        if index >= MAX_PAGES:
            if not found_text:
                raise DocumentError("PDF 前 200 页没有文本层，请先进行 OCR 识别后重试。")
            yield TRUNCATION_NOTE
            return
        text = page.extract_text().strip()
        if text:
            found_text = True
            yield f"[第 {index + 1} 页]\n{text}\n"
        else:
            empty_pages += 1
    if not found_text:
        raise DocumentError("PDF 没有可提取的文本层（可能是扫描件），请先进行 OCR 识别后重试。")
    if empty_pages:
        yield f"[提示：{empty_pages} 页没有文本层，可能包含扫描图片，未进行 OCR。]"


def _office_parts(data: bytes, format: str) -> Generator[str]:
    # Bound expansion before handing an OOXML archive to a format library.
    with ZipFile(BytesIO(data)) as archive:
        if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
            raise DocumentError("文档解压后超过 100MB 限制，请拆分文件后重试。")
    if format == "docx":
        from docx import Document
        from docx.table import Table

        for block in Document(BytesIO(data)).iter_inner_content():
            if isinstance(block, Table):
                for row in block.rows:
                    yield "\t".join(cell.text for cell in row.cells)
            else:
                yield block.text
    elif format == "xlsx":
        from openpyxl import load_workbook

        # Preserve formulas rather than silently dropping cells without caches.
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in workbook:
                yield f"[工作表：{sheet.title}；公式以原始表达式显示]"
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index >= 10000:
                        yield TRUNCATION_NOTE
                        return
                    values = ["" if value is None else str(value) for value in row]
                    if any(values):
                        yield "\t".join(values)
        finally:
            workbook.close()
    elif format == "pptx":
        from pptx import Presentation
        from pptx.shapes.autoshape import Shape
        from pptx.shapes.graphfrm import GraphicFrame
        from pptx.shapes.group import GroupShape

        def shape_parts(shapes: SlideShapes | GroupShapes) -> Iterator[str]:
            for shape in shapes:
                if isinstance(shape, Shape) and shape.has_text_frame:
                    yield shape.text
                if isinstance(shape, GraphicFrame) and shape.has_table:
                    for row in shape.table.rows:
                        yield "\t".join(cell.text for cell in row.cells)
                if isinstance(shape, GroupShape):
                    yield from shape_parts(shape.shapes)

        for index, slide in enumerate(Presentation(BytesIO(data)).slides):
            if index >= MAX_PAGES:
                yield TRUNCATION_NOTE
                return
            yield f"[幻灯片 {index + 1}]"
            yield from shape_parts(slide.shapes)


def parse_document(path: Path, *, max_chars: int = MAX_TEXT_CHARS) -> ParsedDocument:
    """Extract bounded text without Poppler, Office, or network access."""
    if max_chars <= 0:
        raise ValueError("max_chars must be positive")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise DocumentError("资料文件超过 20MB 限制。")
    with path.open("rb") as stream:
        data = stream.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError("资料文件超过 20MB 限制。")
    format = "pdf" if data.startswith(b"%PDF-") else path.suffix.lower().lstrip(".")
    if format not in SUPPORTED_FORMATS:
        raise DocumentError("暂不支持该资料格式，请使用 PDF、DOCX、XLSX 或 PPTX。")
    parts = _pdf_parts(data) if format == "pdf" else _office_parts(data, format)
    text = ""
    truncated = False
    try:
        for part in parts:
            if part == TRUNCATION_NOTE:
                truncated = True
                break
            if not part.strip():
                continue
            addition = ("\n" if text else "") + part
            remaining = max_chars - len(text)
            text += addition[:remaining]
            if len(addition) > remaining:
                truncated = True
                break
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError(f"{format.upper()} 解析失败，文件可能已损坏或加密。") from exc
    finally:
        parts.close()
    if not text.strip():
        raise DocumentError("资料中没有可提取的文本内容。")
    if truncated:
        text += TRUNCATION_NOTE
    return ParsedDocument(text=text, format=format, truncated=truncated)


def main() -> None:
    try:
        if len(sys.argv) != 2:
            raise DocumentError("需要提供一个资料文件路径。")
        if sys.argv[1] == "--office":
            from sztu_code.core.office import execute_office_request

            raw = sys.stdin.buffer.read(1024 * 1024 + 1)
            if len(raw) > 1024 * 1024:
                raise DocumentError("办公请求超过 1MB 限制")
            result = {"result": execute_office_request(json.loads(raw))}
        else:
            result = asdict(parse_document(Path(sys.argv[1])))
    except Exception as exc:
        result = {"error": str(exc)}
    # ASCII JSON escapes make the protocol independent of Windows code pages.
    print(json.dumps(result, ensure_ascii=True))


if __name__ == "__main__":
    main()
