from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from sztu_code.core.config import SztuConfig
from sztu_code.core.office import execute_office_request, office_contract
from sztu_code.core.runner import AgentRunner
from sztu_code.core.task.manager import TaskManager
from sztu_code.core.tools.base import ToolPermission
from sztu_code.core.tools.builtin.office import CreateDocumentTool, ReadDocumentTool


def invoke(root: Path, tool: str, **params: object) -> dict:
    return execute_office_request({"workspace": str(root), "tool": tool, "params": params})


@pytest.mark.parametrize("format,payload,needle,operations", [
    ("docx", {"blocks": [{"kind": "heading", "text": "季度总结"},
                          {"kind": "table", "rows": [["项目", "结果"], ["收入", "100"]]}]},
     "季度总结", [{"kind": "replace_text", "find": "季度总结", "replace": "年度总结"}]),
    ("xlsx", {"sheets": [{"name": "预算", "rows": [["项目", "金额"], ["收入", 100], ["公式", "=B2*2"]]}]},
     "=B2*2", [{"kind": "set_cell", "sheet": "预算", "cell": "B2", "value": 200}]),
    ("pptx", {"slides": [{"title": "季度总结", "bullets": ["收入增加", "下一步计划"], "notes": "讲解备注"}]},
     "讲解备注", [{"kind": "replace_text", "find": "季度总结", "replace": "年度总结"}]),
])
def test_office_roundtrip(root_tmp_path: Path, format: str, payload: dict, needle: str, operations: list) -> None:
    root = root_tmp_path
    source = f"input.{format}"
    made = invoke(root, "create_document", path=source, **payload)
    assert made["verification"] == "saved_and_reopened"
    assert made["visual_verified"] is False
    read = invoke(root, "read_document", path=source)
    assert needle in json.dumps(read, ensure_ascii=False)
    original = (root / source).read_bytes()
    edited = invoke(root, "edit_document", path=f"output.{format}", source=source,
                    source_sha256=read["sha256"], operations=operations)
    assert edited["verification"] == "saved_and_reopened"
    assert (root / source).read_bytes() == original
    reread = invoke(root, "read_document", path=f"output.{format}")
    assert ("200" if format == "xlsx" else "年度总结") in json.dumps(reread, ensure_ascii=False)


@pytest.fixture
def root_tmp_path(tmp_path: Path) -> Path:
    return tmp_path


def test_large_document_pagination_does_not_lose_tail(tmp_path: Path) -> None:
    expected = "甲" * 10000 + "最后的关键信息"
    invoke(tmp_path, "create_document", path="long.docx", blocks=[{"kind": "paragraph", "text": expected}])
    chunks = []
    offset = 0
    while True:
        result = invoke(tmp_path, "read_document", path="long.docx", offset=offset, limit=2)
        chunks.extend(result["blocks"])
        offset = result["next_offset"]
        if offset is None:
            break
    assert "".join(chunk["text"] for chunk in chunks) == expected
    assert [chunk["index"] for chunk in chunks] == list(range(len(chunks)))
    assert all(chunk["location"] == "body/block:1" for chunk in chunks)


def test_xlsx_sheet_coordinate_and_types(tmp_path: Path) -> None:
    invoke(tmp_path, "create_document", path="data.xlsx", sheets=[
        {"name": "第一页", "rows": [["不可混入"]]},
        {"name": "数据", "rows": [["值", "公式"], [123, "=A2*2"], [False, None]]},
    ])
    result = invoke(tmp_path, "read_document", path="data.xlsx", sheet="数据")
    assert result["sheets"] == ["第一页", "数据"]
    assert {block["location"]: block["data_type"] for block in result["blocks"]}["数据!B2"] == "f"
    assert "不可混入" not in json.dumps(result, ensure_ascii=False)
    with pytest.raises(ValueError, match="不存在"):
        invoke(tmp_path, "read_document", path="data.xlsx", sheet="缺失")


@pytest.mark.parametrize("find,replace,text,expected", [
    ("aaa", "X", "aaaaaa", "XX"),
    ("bcde", "中文", "abcdef", "a中文f"),
    ("a", "longer", "abcabc", "longerbclongerbc"),
])
def test_cross_run_replacement_preserves_styles(tmp_path: Path, find: str, replace: str, text: str, expected: str) -> None:
    doc = Document()
    p = doc.add_paragraph()
    for index, char in enumerate(text):
        p.add_run(char).bold = index % 2 == 0
    doc.save(tmp_path / "source.docx")
    invoke(tmp_path, "edit_document", source="source.docx", path="edited.docx", operations=[
        {"kind": "replace_text", "find": find, "replace": replace, "expected_matches": text.count(find)},
    ])
    edited = Document(tmp_path / "edited.docx").paragraphs[0]
    assert edited.text == expected
    assert [run.bold for run in edited.runs] == [index % 2 == 0 for index in range(len(text))]


def test_read_and_edit_headers_and_nested_tables(tmp_path: Path) -> None:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "页眉内容"
    table = doc.add_table(rows=1, cols=1)
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "嵌套内容"
    doc.save(tmp_path / "source.docx")
    result = invoke(tmp_path, "read_document", path="source.docx")
    assert "页眉内容" in json.dumps(result, ensure_ascii=False)
    assert "嵌套内容" in json.dumps(result, ensure_ascii=False)
    invoke(tmp_path, "edit_document", path="out.docx", source="source.docx", operations=[
        {"kind": "replace_text", "find": "嵌套内容", "replace": "新内容"},
    ])
    assert "新内容" in json.dumps(invoke(tmp_path, "read_document", path="out.docx"), ensure_ascii=False)


def test_edit_is_atomic_on_later_failure(tmp_path: Path) -> None:
    invoke(tmp_path, "create_document", path="source.docx", blocks=[{"kind": "paragraph", "text": "original"}])
    source = (tmp_path / "source.docx").read_bytes()
    with pytest.raises(ValueError, match="匹配 0"):
        invoke(tmp_path, "edit_document", source="source.docx", path="out.docx", operations=[
            {"kind": "replace_text", "find": "original", "replace": "changed"},
            {"kind": "replace_text", "find": "missing", "replace": "changed"},
        ])
    assert not (tmp_path / "out.docx").exists()
    assert (tmp_path / "source.docx").read_bytes() == source


def test_no_clobber_and_stale_source(tmp_path: Path) -> None:
    payload = {"path": "source.docx", "blocks": [{"kind": "paragraph", "text": "original"}]}
    invoke(tmp_path, "create_document", **payload)
    original = (tmp_path / "source.docx").read_bytes()
    with pytest.raises(ValueError, match="已存在"):
        invoke(tmp_path, "create_document", **payload)
    assert (tmp_path / "source.docx").read_bytes() == original
    with pytest.raises(ValueError, match="已变化"):
        invoke(tmp_path, "edit_document", source="source.docx", path="out.docx", source_sha256="0" * 64,
               operations=[{"kind": "replace_text", "find": "original", "replace": "new"}])
    assert not (tmp_path / "out.docx").exists()


@pytest.mark.parametrize("path", ["../escape.docx", "C:/escape.docx", "/escape.docx"])
def test_path_boundary(tmp_path: Path, path: str) -> None:
    with pytest.raises(PermissionError):
        invoke(tmp_path, "create_document", path=path, blocks=[{"kind": "paragraph", "text": "x"}])


@pytest.mark.parametrize("name", ["same", "Same", "invalid/name"])
def test_invalid_or_duplicate_sheet_names(tmp_path: Path, name: str) -> None:
    with pytest.raises(ValueError, match="非法或重复"):
        invoke(tmp_path, "create_document", path="bad.xlsx", sheets=[
            {"name": "same", "rows": [[1]]}, {"name": name, "rows": [[2]]},
        ])
    assert not (tmp_path / "bad.xlsx").exists()


async def test_permissions_and_model_visible_tools(tmp_path: Path) -> None:
    assert ReadDocumentTool(tmp_path).required_permission == ToolPermission.READ_ONLY
    writer = CreateDocumentTool(tmp_path, ["reports/**"])
    assert writer.classify_permission({"path": "elsewhere.docx"}) == ToolPermission.DANGER_FULL_ACCESS
    with pytest.raises(PermissionError):
        await writer.invoke({"path": "elsewhere.docx", "blocks": [{"kind": "paragraph", "text": "x"}]})
    runner = AgentRunner(SztuConfig())
    registry = runner._build_registry(TaskManager(tmp_path / "tasks"), workspace_root=tmp_path)
    schemas = {item["name"]: item for item in registry.tool_schemas()}
    assert {"read_document", "create_document", "edit_document"} <= schemas.keys()
    assert "DOCX" in schemas["read_file"]["description"]
    result = await registry.get("create_document").invoke({
        "description": "创建报告", "path": "agent.docx", "blocks": [{"kind": "paragraph", "text": "模型工具调用"}],
    })
    assert not result.is_error, result.content
    assert (tmp_path / "agent.docx").exists()


def test_typescript_schema_matches_python() -> None:
    path = Path(__file__).resolve().parents[3] / "packages/runtime-ts/src/office-contract.ts"
    text = path.read_text(encoding="utf-8")
    contract = json.loads(text.split("export const officeToolContracts = ", 1)[1].rstrip(";\n"))
    assert contract == office_contract()
