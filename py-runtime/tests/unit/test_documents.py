from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from sztu_code.core.documents import (
    MAX_FILE_BYTES,
    TRUNCATION_NOTE,
    DocumentError,
    parse_document,
)
from sztu_code.core.tools.builtin.read_file import ReadFileTool


def write_pdf(path: Path, texts: list[str], password: str | None = None) -> None:
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=612, height=792)
        chars = sorted(set(text))
        mapping = {char: index + 1 for index, char in enumerate(chars)}
        cmap = DecodedStreamObject()
        cmap.set_data((
            "/CIDInit /ProcSet findresource begin 12 dict begin begincmap\n"
            "/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n"
            "/CMapName /Test def /CMapType 2 def\n"
            "1 begincodespacerange <00> <FF> endcodespacerange\n"
            f"{len(chars)} beginbfchar\n"
            + "\n".join(f"<{code:02X}> <{char.encode('utf-16-be').hex()}>" for char, code in mapping.items())
            + "\nendbfchar endcmap CMapName currentdict /CMap defineresource pop end end"
        ).encode())
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
            NameObject("/ToUnicode"): writer._add_object(cmap),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})
        })
        content = DecodedStreamObject()
        content.set_data(f"BT /F1 12 Tf 72 720 Td <{''.join(f'{mapping[c]:02X}' for c in text)}> Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(content)
    if password is not None:
        writer.encrypt(password, algorithm="AES-256")
    writer.write(path)


def test_pdf_chinese_pages_and_no_poppler(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "学习 资料.PDF"
    write_pdf(path, ["中文资料解析 Hello", "第二页 内容"])
    monkeypatch.setenv("PATH", "")
    result = parse_document(path)
    assert "中文资料解析 Hello" in result.text
    assert "[第 2 页]" in result.text
    assert "第二页 内容" in result.text
    assert not result.truncated


def test_pdf_signature_without_extension_and_size(tmp_path: Path) -> None:
    path = tmp_path / "document"
    write_pdf(path, ["Hello"])
    # A PDF larger than the text attachment limit is still a document.
    with path.open("ab") as stream:
        stream.write(b" " * (1024 * 1024))
    assert parse_document(path).format == "pdf"


@pytest.mark.parametrize("password", ["secret", ""])
def test_pdf_encryption(tmp_path: Path, password: str) -> None:
    path = tmp_path / "encrypted.pdf"
    write_pdf(path, ["content"], password)
    if password:
        with pytest.raises(DocumentError, match="已加密"):
            parse_document(path)
    else:
        assert "content" in parse_document(path).text


def test_scanned_and_mixed_pdf(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    write_pdf(path, [""])
    with pytest.raises(DocumentError, match="OCR"):
        parse_document(path)
    write_pdf(path, ["正文", ""])
    assert "1 页没有文本层" in parse_document(path).text


def test_truncation_and_page_limit(tmp_path: Path) -> None:
    path = tmp_path / "long.pdf"
    write_pdf(path, ["A" * 100])
    result = parse_document(path, max_chars=20)
    assert result.truncated and result.text.endswith(TRUNCATION_NOTE)
    assert len(result.text) == 20 + len(TRUNCATION_NOTE)
    write_pdf(path, ["A"] * 201)
    result = parse_document(path)
    assert result.truncated and "[第 201 页]" not in result.text


@pytest.mark.parametrize("extension", ["pdf", "docx", "xlsx", "pptx"])
def test_corrupt_and_oversized(tmp_path: Path, extension: str) -> None:
    path = tmp_path / f"broken.{extension}"
    path.write_bytes(b"not a document")
    with pytest.raises(DocumentError, match="解析失败"):
        parse_document(path)
    with path.open("wb") as stream:
        stream.truncate(MAX_FILE_BYTES + 1)
    with pytest.raises(DocumentError, match="20MB"):
        parse_document(path)


def test_docx_preserves_paragraph_table_order(tmp_path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("资料开头")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "金额"
    doc.add_paragraph("资料结尾")
    path = tmp_path / "资料.docx"
    doc.save(path)
    text = parse_document(path).text
    assert text.index("资料开头") < text.index("项目\t金额") < text.index("资料结尾")


def test_xlsx_preserves_formulas_and_sheet_names(tmp_path: Path) -> None:
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "预算"
    book.active.append(["费用", 123, "=B1*2"])
    book.create_sheet("说明").append(["第二表"])
    path = tmp_path / "资料.xlsx"
    book.save(path)
    text = parse_document(path).text
    assert "预算" in text and "费用\t123\t=B1*2" in text and "第二表" in text


def test_pptx_text_table_and_group(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "分组标题"
    table = slide.shapes.add_table(1, 2, 0, 0, Inches(3), Inches(1)).table
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "内容"
    path = tmp_path / "资料.pptx"
    slides.save(path)
    text = parse_document(path).text
    assert "幻灯片 1" in text and "分组标题" in text and "项目\t内容" in text


def test_archive_expansion_limit(tmp_path: Path) -> None:
    path = tmp_path / "large.docx"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", b"x" * (101 * 1024 * 1024))
    with pytest.raises(DocumentError, match="100MB"):
        parse_document(path)


async def test_read_file_extracts_document_and_reports_errors(tmp_path: Path) -> None:
    path = tmp_path / "input.pdf"
    write_pdf(path, ["可供聊天引用的内容"])
    result = await ReadFileTool(tmp_path).invoke({"path": path.name})
    assert not result.is_error and "可供聊天引用的内容" in result.content
    path.write_bytes(b"broken")
    result = await ReadFileTool(tmp_path).invoke({"path": path.name})
    assert result.is_error and "解析失败" in result.content


def test_desktop_cli_protocol_is_unicode_safe(tmp_path: Path) -> None:
    import sztu_code.core.documents as documents

    path = tmp_path / "中文 路径.pdf"
    write_pdf(path, ["资料内容"])
    result = subprocess.run(
        [sys.executable, documents.__file__, str(path)], capture_output=True,
        env={**os.environ, "PATH": "", "PYTHONIOENCODING": "ascii"}, check=True,
    )
    assert "资料内容" in json.loads(result.stdout)["text"]


@pytest.mark.skipif(not os.environ.get("DOCUMENT_PARSER_EXECUTABLE"), reason="requires desktop bundle")
def test_bundled_parser_without_system_python_or_poppler(tmp_path: Path) -> None:
    write_pdf(tmp_path / "资料.pdf", ["中文资料"], password="")
    test_docx_preserves_paragraph_table_order(tmp_path)
    test_xlsx_preserves_formulas_and_sheet_names(tmp_path)
    test_pptx_text_table_and_group(tmp_path)
    for path in tmp_path.iterdir():
        result = subprocess.run(
            [os.environ["DOCUMENT_PARSER_EXECUTABLE"], str(path)], capture_output=True,
            env={**os.environ, "PATH": "", "PYTHONHOME": "", "PYTHONPATH": ""},
            check=True, timeout=30,
        )
        assert json.loads(result.stdout)["text"] == parse_document(path).text
